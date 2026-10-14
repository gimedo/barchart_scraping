"""
Descarga cotizaciones de futuros desde Barchart y las guarda con el formato:

symbol | last_price | price_change | open_price | high_price | low_price |
previous_price | volume | open_interest | fecha_reporte | fecha_hora_obtencion |
mes_anio | contract | contract_month | contract_year | mes_adelante | mes

Instalacion (una sola vez):
    pip install pandas playwright openpyxl
    playwright install chromium

Uso:
    python barchart_scraper.py

Barchart protege el sitio con un desafio JavaScript (HTTP 202), asi que se usa
un navegador real (Playwright) para pasarlo y luego se llama al endpoint
interno desde dentro del navegador.

Salidas (en la carpeta actual):
    barchart_ultima.xlsx / .csv      -> ultima captura
    barchart_historico.xlsx / .csv   -> acumulado de todas las capturas
    (cierra el .xlsx en Excel antes de ejecutar, o Windows lo bloqueara)

AVISO: usa un endpoint interno no documentado. Puede cambiar o bloquearte.
Revisa los terminos de uso de Barchart antes de automatizarlo.
"""
import json
import os
import re
import sys
import time
from datetime import datetime
from urllib.parse import unquote, urlencode

import pandas as pd
from playwright.sync_api import sync_playwright

# --------------------------------------------------------------------------
# Configuracion
# --------------------------------------------------------------------------
# Raices de contratos a descargar. Antes tenias CL, RB, HO y QA.
# Para granos/indice dolar agrega: "ZC", "ZS", "ZM", "ZL", "ZW", "KE", "DX"
ROOTS = ["CL", "RB", "HO", "QA"]

# Desfase de mes_adelante por raiz (deducido de tu muestra: RB y HO = +1).
# Verificalo contra tus datos viejos y ajustalo si hace falta.
MES_ADELANTE_OFFSET = {"RB": 1, "HO": 1}

# False = abre una ventana de Chrome (pasa mejor el anti-bot). True = sin ventana.
HEADLESS = os.environ.get("BARCHART_HEADLESS", "false").lower() in ("1", "true", "yes", "si")
CDP_URL = os.environ.get("BARCHART_CDP_URL", "")   # p. ej. http://127.0.0.1:9222 (Chrome ya abierto)
ESPERA_DESAFIO_SEG = 25

GUARDAR_EXCEL = os.environ.get("BARCHART_EXCEL", "true").lower() in ("1", "true", "yes", "si")
GUARDAR_CSV = True
OUT_HISTORICO = "barchart_historico"   # se agrega .xlsx / .csv
OUT_ULTIMA = "barchart_ultima"

BASE = "https://www.barchart.com"
API = BASE + "/proxies/core-api/v1/quotes/get"
PAGE_SIZE = 100

FIELDS = (
    "symbol,contractSymbol,lastPrice,priceChange,openPrice,highPrice,lowPrice,"
    "previousPrice,volume,openInterest,tradeTime,symbolCode,symbolType,hasOptions"
)

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)

MONTH_CODE = {
    "F": "Jan", "G": "Feb", "H": "Mar", "J": "Apr", "K": "May", "M": "Jun",
    "N": "Jul", "Q": "Aug", "U": "Sep", "V": "Oct", "X": "Nov", "Z": "Dec",
}
MONTH_NUM = {k: i + 1 for i, k in enumerate("FGHJKMNQUVXZ")}
MES_ES = {
    "Jan": "ene", "Feb": "feb", "Mar": "mar", "Apr": "abr", "May": "may",
    "Jun": "jun", "Jul": "jul", "Aug": "ago", "Sep": "sep", "Oct": "oct",
    "Nov": "nov", "Dec": "dic",
}
SYMBOL_RE = re.compile(r"^(.+?)([FGHJKMNQUVXZ])(\d{2})$")


# --------------------------------------------------------------------------
# Descarga (Playwright)
# --------------------------------------------------------------------------
JS_FETCH = """
async ({url, token}) => {
  const headers = {'accept': 'application/json'};
  if (token) headers['x-xsrf-token'] = token;
  const r = await fetch(url, {credentials: 'same-origin', headers});
  return {status: r.status, text: await r.text()};
}
"""

CAPTURA = {}  # token visto en las llamadas que hace la propia pagina
SIN_TOKEN_OK = False  # pasa a True si una raiz ya funciono sin token


def on_request(req):
    if "/proxies/core-api/" in req.url:
        tok = req.headers.get("x-xsrf-token")
        if tok:
            CAPTURA["token"] = unquote(tok)


def esperar_token(context, timeout=ESPERA_DESAFIO_SEG):
    """Busca el token en la cookie XSRF-TOKEN o en las peticiones de la pagina."""
    fin = time.time() + timeout
    while time.time() < fin:
        for c in context.cookies():
            if c["name"].upper() in ("XSRF-TOKEN", "X-XSRF-TOKEN"):
                return unquote(c["value"])
        if CAPTURA.get("token"):
            return CAPTURA["token"]
        time.sleep(1)
    return None


def fetch_root(page, context, root: str) -> list[dict]:
    """Descarga todos los contratos de una raiz, paginando."""
    page.goto(f"{BASE}/futures/quotes/{root}*0/futures-prices",
              wait_until="domcontentloaded", timeout=60000)
    global SIN_TOKEN_OK
    token = None if SIN_TOKEN_OK else esperar_token(context)
    if not token and not SIN_TOKEN_OK:
        raros = [c["name"] for c in context.cookies()
                 if any(k in c["name"].lower() for k in ("xsrf", "csrf", "token", "session", "waf"))]
        print(f"  {root}: sin token XSRF; se intenta sin el. Cookies relevantes: {raros}")

    rows: list[dict] = []
    pag = 1
    while True:
        params = {
            "fields": FIELDS,
            "lists": "futures.contractInRoot",
            "root": root,
            "meta": "field.shortName,field.type,field.description,lists.lastUpdate",
            "hasOptions": "true",
            "page": pag,
            "limit": PAGE_SIZE,
            "raw": "1",
        }
        url = f"{API}?{urlencode(params)}"
        res = page.evaluate(JS_FETCH, {"url": url, "token": token})
        if res["status"] != 200:
            raise RuntimeError(f"{root} pagina {pag}: HTTP {res['status']} {res['text'][:200]!r}")

        payload = json.loads(res["text"])
        data = payload.get("data", [])
        for item in data:
            rows.append(item.get("raw", item))

        total = payload.get("total") or payload.get("count")
        if len(data) < PAGE_SIZE or (total and len(rows) >= int(total)):
            break
        pag += 1
        time.sleep(1)
    if not token:
        SIN_TOKEN_OK = True  # las siguientes raices no esperan el token
    return rows


# --------------------------------------------------------------------------
# Transformacion al formato historico
# --------------------------------------------------------------------------
def to_number(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def build_dataframe(raw_rows: list[dict], ahora: datetime) -> pd.DataFrame:
    records = []
    fechas = []

    for r in raw_rows:
        symbol = r.get("symbol")
        m = SYMBOL_RE.match(symbol or "")
        if not m:  # descarta simbolos tipo CL*0, etc.
            continue
        root, code, yy = m.groups()

        trade = r.get("tradeTime")
        fecha = None
        if trade:
            try:
                if isinstance(trade, (int, float)):
                    fecha = datetime.fromtimestamp(trade).date()
                else:
                    fecha = pd.to_datetime(trade).date()
            except Exception:
                fecha = None
        if fecha:
            fechas.append(fecha)

        year = 2000 + int(yy)
        meses_dif = (year - ahora.year) * 12 + MONTH_NUM[code] - ahora.month
        mes_en = MONTH_CODE[code]

        records.append({
            "symbol": symbol,
            "last_price": to_number(r.get("lastPrice")),
            "price_change": to_number(r.get("priceChange")),
            "open_price": to_number(r.get("openPrice")) or 0.0,
            "high_price": to_number(r.get("highPrice")),
            "low_price": to_number(r.get("lowPrice")),
            "previous_price": to_number(r.get("previousPrice")),
            "volume": to_number(r.get("volume")) or 0.0,
            "open_interest": to_number(r.get("openInterest")) or 0.0,
            "_fecha": fecha,
            "mes_anio": f"{mes_en} '{yy}",
            "contract": root,
            "contract_month": code,
            "contract_year": yy,
            "mes_adelante": meses_dif + MES_ADELANTE_OFFSET.get(root, 0),
            "mes": MES_ES[mes_en],
        })

    df = pd.DataFrame(records)
    if df.empty:
        return df

    # fecha_reporte: dia de negociacion mas reciente devuelto por la API
    fecha_rep = max(fechas) if fechas else ahora.date()
    df["fecha_reporte"] = fecha_rep.isoformat()
    df["fecha_hora_obtencion"] = ahora.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
    df = df.drop(columns=["_fecha"])

    cols = [
        "symbol", "last_price", "price_change", "open_price", "high_price",
        "low_price", "previous_price", "volume", "open_interest",
        "fecha_reporte", "fecha_hora_obtencion", "mes_anio", "contract",
        "contract_month", "contract_year", "mes_adelante", "mes",
    ]
    return df[cols]


# --------------------------------------------------------------------------
# Salida (CSV / Excel)
# --------------------------------------------------------------------------
COLUMNAS = [
    "symbol", "last_price", "price_change", "open_price", "high_price",
    "low_price", "previous_price", "volume", "open_interest",
    "fecha_reporte", "fecha_hora_obtencion", "mes_anio", "contract",
    "contract_month", "contract_year", "mes_adelante", "mes",
]


def para_excel(df: pd.DataFrame) -> pd.DataFrame:
    """Convierte las columnas de fecha a tipos reales de fecha para Excel."""
    d = df.copy()
    d["fecha_reporte"] = pd.to_datetime(d["fecha_reporte"]).dt.date
    d["fecha_hora_obtencion"] = pd.to_datetime(d["fecha_hora_obtencion"])
    for c in ("symbol", "mes_anio", "contract", "contract_month", "contract_year", "mes"):
        d[c] = d[c].astype(str)
    return d[COLUMNAS]


def escribir_excel(df: pd.DataFrame, ruta: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="datos", index=False)
        ws = xw.sheets["datos"]

        base = Font(name="Arial", size=10)
        head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="1F3864")
        for row in ws.iter_rows():
            for cell in row:
                cell.font = base
        for cell in ws[1]:
            cell.font = head
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")

        formatos = {
            "volume": "#,##0", "open_interest": "#,##0",
            "fecha_reporte": "yyyy-mm-dd",
            "fecha_hora_obtencion": "yyyy-mm-dd hh:mm:ss.000",
        }
        for idx, nombre in enumerate(COLUMNAS, start=1):
            letra = get_column_letter(idx)
            ancho = max(len(nombre), 12)
            if nombre == "fecha_hora_obtencion":
                ancho = 24
            ws.column_dimensions[letra].width = ancho + 2
            if nombre in formatos:
                for cell in ws[letra][1:]:
                    cell.number_format = formatos[nombre]

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions


def guardar(df: pd.DataFrame) -> None:
    # --- CSV
    if GUARDAR_CSV:
        df.to_csv(OUT_ULTIMA + ".csv", index=False)
        existe = os.path.exists(OUT_HISTORICO + ".csv")
        df.to_csv(OUT_HISTORICO + ".csv", mode="a", header=not existe, index=False)
        print(f"CSV   -> {OUT_ULTIMA}.csv y {OUT_HISTORICO}.csv")

    # --- Excel
    if GUARDAR_EXCEL:
        nuevo = para_excel(df)
        try:
            escribir_excel(nuevo, OUT_ULTIMA + ".xlsx")
            ruta_h = OUT_HISTORICO + ".xlsx"
            if os.path.exists(ruta_h):
                previo = pd.read_excel(ruta_h, sheet_name="datos", dtype={
                    "symbol": str, "mes_anio": str, "contract": str,
                    "contract_month": str, "contract_year": str, "mes": str})
                historico = pd.concat([para_excel(previo), nuevo], ignore_index=True)
            else:
                historico = nuevo
            escribir_excel(historico, ruta_h)
            print(f"Excel -> {OUT_ULTIMA}.xlsx y {ruta_h} ({len(historico)} filas acumuladas)")
        except PermissionError:
            print("ERROR: no se pudo escribir el Excel. Cierralo si lo tienes abierto.",
                  file=sys.stderr)


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def main() -> int:
    ahora = datetime.now()
    raw_all: list[dict] = []

    with sync_playwright() as pw:
        if CDP_URL:  # conectarse a un Chrome ya abierto como servicio
            browser = pw.chromium.connect_over_cdp(CDP_URL)
            context = browser.contexts[0] if browser.contexts else browser.new_context()
        else:
            browser = pw.chromium.launch(headless=HEADLESS)
            context = browser.new_context(
                locale="en-US",
                viewport={"width": 1366, "height": 850},
            )
        page = context.new_page()
        page.on("request", on_request)
        for root in ROOTS:
            try:
                filas = fetch_root(page, context, root)
                print(f"{root}: {len(filas)} contratos")
                raw_all.extend(filas)
            except Exception as e:
                print(f"{root}: ERROR {e}", file=sys.stderr)
            time.sleep(2)
        page.close()
        browser.close()

    df = build_dataframe(raw_all, ahora)
    if df.empty:
        print("Sin datos; no se escribio nada.", file=sys.stderr)
        return 1

    guardar(df)
    print(f"{len(df)} filas descargadas")
    print(df.head())
    return 0


if __name__ == "__main__":
    sys.exit(main())
